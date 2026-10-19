"""
Langkah 4 - Membuat workbook Excel "Analisis Turnover Karyawan.xlsx".

Semua angka di workbook ditulis sebagai rumus Excel (COUNTIFS, SUMIFS, PERCENTILE,
MEDIAN), jadi akan dihitung ulang otomatis saat file dibuka di Excel.
Butuh paket openpyxl.
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
ROOT = Path(__file__).resolve().parents[1]
d = pd.read_csv(ROOT / 'data' / 'ibm_hr_attrition.csv', encoding='utf-8-sig')
cols={'EmployeeNumber':'ID Karyawan','Age':'Usia','Attrition':'Resign','Department':'Departemen','JobRole':'Jabatan','JobLevel':'Level Jabatan','OverTime':'Lembur','MaritalStatus':'Status Pernikahan','BusinessTravel':'Dinas Luar','MonthlyIncome':'Gaji Bulanan (USD)','YearsAtCompany':'Masa Kerja (tahun)','JobSatisfaction':'Kepuasan Kerja (1-4)','WorkLifeBalance':'Work-Life Balance (1-4)','StockOptionLevel':'Level Opsi Saham (0-3)'}
d=d[list(cols)].rename(columns=cols)
tr={'Research & Development':'Riset & Pengembangan','Sales':'Penjualan','Human Resources':'SDM'}
d['Departemen']=d['Departemen'].map(tr); d['Resign']=d['Resign'].map({'Yes':'Ya','No':'Tidak'}); d['Lembur']=d['Lembur'].map({'Yes':'Ya','No':'Tidak'})
d['Status Pernikahan']=d['Status Pernikahan'].map({'Single':'Lajang','Married':'Menikah','Divorced':'Cerai'})
d['Dinas Luar']=d['Dinas Luar'].map({'Non-Travel':'Tidak pernah','Travel_Rarely':'Jarang','Travel_Frequently':'Sering'})
F='Arial'; f=lambda **k: Font(name=F,**k)
hdr=PatternFill('solid',start_color='1F3864'); inp=PatternFill('solid',start_color='FFFF00'); thin=Side(style='thin',color='BFBFBF')
wb=Workbook(); R=wb.active; R.title='Ringkasan'; A=wb.create_sheet('Analisis'); D=wb.create_sheet('Data'); S=wb.create_sheet('Sumber')
# Data sheet
heads=list(d.columns)+['Keluar (1/0)','Kelompok Masa Kerja','Kuartil Gaji']
for j,h in enumerate(heads,1):
    c=D.cell(1,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='center')
n=len(d)
for i,row in enumerate(d.itertuples(index=False),2):
    for j,v in enumerate(row,1): D.cell(i,j,v.item() if hasattr(v,'item') else v).font=f()
    D.cell(i,15,f'=IF(C{i}="Ya",1,0)').font=f()
    D.cell(i,16,f'=IF(K{i}<1,"a. <1 tahun",IF(K{i}<2,"b. 1 tahun",IF(K{i}<3,"c. 2 tahun",IF(K{i}<=5,"d. 3-5 tahun",IF(K{i}<=10,"e. 6-10 tahun","f. >10 tahun")))))').font=f()
    D.cell(i,17,f'=IF(J{i}<=PERCENTILE($J$2:$J${n+1},0.25),"Q1 (terendah)",IF(J{i}<=PERCENTILE($J$2:$J${n+1},0.5),"Q2",IF(J{i}<=PERCENTILE($J$2:$J${n+1},0.75),"Q3","Q4 (tertinggi)")))').font=f()
for j,w in enumerate([11,7,8,20,24,9,8,11,12,12,11,11,11,11,9,16,14],1): D.column_dimensions[D.cell(1,j).column_letter].width=w
D.freeze_panes='A2'; D.auto_filter.ref=f'A1:Q{n+1}'
L=n+1
# Analisis sheet
A['A1']='Analisis turnover per dimensi'; A['A1'].font=f(bold=True,size=14)
A['A2']='Semua angka dihitung dengan rumus COUNTIFS dari sheet Data.'; A['A2'].font=f(italic=True,color='595959')
blocks=[('Masa kerja','P',['a. <1 tahun','b. 1 tahun','c. 2 tahun','d. 3-5 tahun','e. 6-10 tahun','f. >10 tahun']),
 ('Lembur','G',['Ya','Tidak']),('Departemen','D',['Penjualan','SDM','Riset & Pengembangan']),
 ('Level jabatan','F',[1,2,3,4,5]),('Kuartil gaji','Q',['Q1 (terendah)','Q2','Q3','Q4 (tertinggi)']),
 ('Status pernikahan','H',['Lajang','Menikah','Cerai']),('Dinas luar','I',['Sering','Jarang','Tidak pernah']),
 ('Jabatan','E',sorted(d['Jabatan'].unique()))]
r=4; refs={}
for title,col,vals in blocks:
    for j,h in enumerate([title,'Karyawan','Resign','Tingkat turnover'],1):
        c=A.cell(r,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=hdr
    top=r+1
    for k,v in enumerate(vals):
        rr=top+k; A.cell(rr,1,v).font=f()
        A.cell(rr,2,f'=COUNTIFS(Data!${col}$2:${col}${L},A{rr})').font=f()
        A.cell(rr,3,f'=COUNTIFS(Data!${col}$2:${col}${L},A{rr},Data!$C$2:$C${L},"Ya")').font=f()
        c=A.cell(rr,4,f'=IFERROR(C{rr}/B{rr},0)'); c.font=f(); c.number_format='0.0%'
        for j in range(1,5): A.cell(rr,j).border=Border(bottom=thin)
    refs[title]=(top,top+len(vals)-1)
    ch=BarChart(); ch.type='bar'; ch.title=f'Tingkat turnover per {title.lower()}'; ch.legend=None; ch.style=10
    ch.add_data(Reference(A,min_col=4,min_row=top,max_row=top+len(vals)-1)); ch.set_categories(Reference(A,min_col=1,min_row=top,max_row=top+len(vals)-1))
    ch.y_axis.numFmt='0%'; ch.y_axis.majorGridlines=None; ch.x_axis.scaling.orientation='maxMin'; ch.y_axis.delete=False; ch.x_axis.delete=False
    ch.dataLabels=DataLabelList(); ch.dataLabels.showVal=True; ch.dataLabels.showCatName=False; ch.dataLabels.showSerName=False; ch.dataLabels.showLegendKey=False; ch.dataLabels.numFmt='0%'
    ch.series[0].graphicalProperties.solidFill='2A78D6'; ch.height=max(4.5,0.9*len(vals)+2); ch.width=13
    A.add_chart(ch,f'F{r}')
    r=max(top+len(vals), r+int(ch.height/0.5)+1)+2
A.column_dimensions['A'].width=26
for c in 'BC': A.column_dimensions[c].width=12
A.column_dimensions['D'].width=17
# Ringkasan
R['A1']='Analisis Turnover Karyawan'; R['A1'].font=f(bold=True,size=16)
R['A2']='Dataset contoh IBM HR Analytics (1.470 karyawan, data fiktif buatan IBM). Penulis: Ratu Adisya Famela.'; R['A2'].font=f(italic=True,color='595959')
R['A4']='Indikator utama'; R['A4'].font=f(bold=True,size=12)
kp=[('Jumlah karyawan',f'=COUNTA(Data!A2:A{L})','#,##0'),('Karyawan resign',f'=SUM(Data!O2:O{L})','#,##0'),('Tingkat turnover','=B6/B5','0.0%'),
('Resign dengan masa kerja < 2 tahun',f'=COUNTIFS(Data!K2:K{L},"<2",Data!C2:C{L},"Ya")','#,##0'),('Porsi resign dalam 2 tahun pertama','=B8/B6','0.0%'),
('Median gaji bulanan, yang resign (USD)',f'=MEDIAN(IF(Data!C2:C{L}="Ya",Data!J2:J{L}))','#,##0'),('Median gaji bulanan, yang bertahan (USD)',f'=MEDIAN(IF(Data!C2:C{L}="Tidak",Data!J2:J{L}))','#,##0')]
for k,(a,b,nf) in enumerate(kp,5):
    R.cell(k,1,a).font=f(); c=R.cell(k,2,b); c.font=f(bold=True); c.number_format=nf
R['B10'].value=None; R['B11'].value=None
from openpyxl.worksheet.formula import ArrayFormula
R['B10']=ArrayFormula('B10',f'=MEDIAN(IF(Data!C2:C{L}="Ya",Data!J2:J{L}))'); R['B11']=ArrayFormula('B11',f'=MEDIAN(IF(Data!C2:C{L}="Tidak",Data!J2:J{L}))')
for c in ('B10','B11'): R[c].font=f(bold=True); R[c].number_format='#,##0'
R['A13']='Estimasi biaya penggantian karyawan'; R['A13'].font=f(bold=True,size=12)
R['A14']='Faktor biaya terendah (× gaji tahunan)'; R['B14']=0.5
R['A15']='Faktor biaya tertinggi (× gaji tahunan)'; R['B15']=2
for c in ('B14','B15'): R[c].font=f(color='0000FF'); R[c].fill=inp; R[c].number_format='0.0"x"'
R['C14']='Sumber: Gallup (2019), biaya mengganti karyawan 0,5–2× gaji tahunan. Ubah sel kuning untuk skenario lain.'; R['C14'].font=f(italic=True,color='595959',size=9)
R['A16']='Total gaji tahunan karyawan yang resign (USD)'; R['B16']=f'=SUMIFS(Data!J2:J{L},Data!C2:C{L},"Ya")*12'
R['A17']='Estimasi biaya penggantian, terendah (USD)'; R['B17']='=B16*B14'
R['A18']='Estimasi biaya penggantian, tertinggi (USD)'; R['B18']='=B16*B15'
for k in range(14,19): R.cell(k,1).font=f()
for c in ('B16','B17','B18'): R[c].font=f(bold=True); R[c].number_format='$#,##0'
R['A20']='Temuan'; R['A20'].font=f(bold=True,size=12)
fs=['Turnover tahun pertama 34-36%, lebih dari 2x rata-rata (16,1%); 32% dari semua yang resign keluar dalam 2 tahun pertama.',
'Karyawan yang lembur resign 30,5%, hampir 3x yang tidak lembur (10,4%). Level 1 + lembur: 53%.',
'Level jabatan 1 (26,3%) dan kuartil gaji terendah (29,3%) paling rawan; median gaji yang resign USD 3.202 vs 5.204.',
'Sales Representative (39,8%) dan Laboratory Technician (23,9%) adalah jabatan paling rawan resign.',
'Karyawan lajang (25,5%) dan yang sering dinas luar (24,9%) lebih sering resign.']
for k,t in enumerate(fs,21): R.cell(k,1,f'{k-20}. {t}').font=f()
R['A27']='Rekomendasi'; R['A27'].font=f(bold=True,size=12)
rs=['Perkuat onboarding dan mentoring 12 bulan pertama, dengan evaluasi rutin di bulan ke-3, 6, dan 12.',
'Batasi dan pantau jam lembur per tim, terutama untuk level jabatan 1.',
'Tinjau gaji level awal terhadap pasar untuk jabatan dengan turnover tertinggi.',
'Lakukan stay interview dan exit interview terstruktur untuk Sales Representative dan Laboratory Technician.']
for k,t in enumerate(rs,28): R.cell(k,1,f'{k-27}. {t}').font=f()
R['A33']='Catatan: dataset tidak memuat data kehadiran (absensi), sehingga analisis kehadiran tidak dilakukan.'; R['A33'].font=f(italic=True,color='595959',size=9)
R.column_dimensions['A'].width=46; R.column_dimensions['B'].width=16
# Sumber
src=[('Dataset','IBM HR Analytics Employee Attrition (data fiktif buatan IBM)'),('Diambil dari','https://github.com/IBM/employee-attrition-aif360 (lisensi Apache 2.0)'),
('Biaya penggantian','Gallup (13 Maret 2019), "This Fixable Problem Costs U.S. Businesses $1 Trillion": https://www.gallup.com/workplace/247391/fixable-problem-costs-businesses-trillion.aspx'),
('Mata uang','Gaji dalam dataset berupa angka tanpa negara; diperlakukan sebagai USD.')]
for k,(a,b) in enumerate(src,1): S.cell(k,1,a).font=f(bold=True); S.cell(k,2,b).font=f()
S.column_dimensions['A'].width=20; S.column_dimensions['B'].width=110
wb.calculation.fullCalcOnLoad = True  # Excel menghitung semua rumus saat file dibuka
wb.save(ROOT / 'Analisis Turnover Karyawan.xlsx')
print('Tersimpan -> Analisis Turnover Karyawan.xlsx')
